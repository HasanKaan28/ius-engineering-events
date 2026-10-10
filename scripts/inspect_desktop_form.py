import os
import docx

docx_path = r'C:\Users\Kaan\Desktop\IEC\01_Resmi_Basvuru_Formlari\student_club_registration_form_f252.docx'
if not os.path.exists(docx_path):
    docx_path = r'C:\Users\Kaan\Desktop\IEC\student_club_registration_form_f252.docx'
if not os.path.exists(docx_path):
    docx_path = r'C:\Users\Kaan\Desktop\student_club_registration_form_f252.docx'

doc = docx.Document(docx_path)
with open(r'scripts\inspect_desktop_form.txt', 'w', encoding='utf-8') as f:
    f.write(f'PARAGRAPHS COUNT: {len(doc.paragraphs)}\n')
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            f.write(f'P[{i}]: {p.text.strip()}\n')

    f.write(f'\nTABLES COUNT: {len(doc.tables)}\n')
    for t_idx, t in enumerate(doc.tables):
        f.write(f'\n=== TABLE {t_idx} (rows: {len(t.rows)}, cols: {len(t.columns)}) ===\n')
        for r_idx, row in enumerate(t.rows):
            row_vals = []
            for c_idx, cell in enumerate(row.cells):
                txt = cell.text.strip().replace('\n', ' [NL] ')
                row_vals.append(f'C{c_idx}: "{txt}"')
            f.write(f'  R[{r_idx}]: {" | ".join(row_vals)}\n')

print('Inspection saved to scripts/inspect_desktop_form.txt')
